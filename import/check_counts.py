import openpyxl
import os

path = 'import_validation.xlsx'
if os.path.exists(path):
    wb = openpyxl.load_workbook(path)
    for sheet in wb.sheetnames:
        count = wb[sheet].max_row - 1
        print(f"{sheet}: {count} entries")
else:
    print(f"{path} not found")
