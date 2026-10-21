import json
import logging
import time
from typing import Any

from openpyxl import load_workbook
from rich.logging import RichHandler
from sqlmodel import Session, select

from utilities.config import DATA_ROOT, DB_STRING, EXCEL
from utilities.db import create_all, engine, get_or_create
from utilities.model import (
    CollectionYear,
    Consultant,
    Coordinate,
    Country,
    DataCollector,
    Fieldwork,
    FieldworkConsultant,
    FieldworkDataCollector,
    Location,
    Place,
    Project,
    Reflex,
    Region,
    Source,
    SourceEntries,
    Variety,
)

from .blr import import_blr, reconstruction_ids
from .columns import (
    COLLECTOR_PATTERN,
    CONSULTANT_PATTERN,
    EXPECTED_COLUMNS,
    SOURCE_GROUPS,
    WHITESPACE_CHECK_FIELDS,
)
from .entries import (
    CATEGORY_CLASSES,
    LEVEL_ORDER,
    BlrIpvResolvedEntry,
    BlrTextToNoteEntry,
    BlrUnresolvedLinkEntry,
    CoordinateConflictEntry,
    CriticalErrorEntry,
    DuplicateBeginningIdEntry,
    FallbackUsedEntry,
    MissingColumnEntry,
    RowContext,
    SkippedNoTermEntry,
    SourceIncompleteEntry,
    SourceLocation,
    SuspiciousDuplicateEntry,
    UncertaintyEntry,
    WhitespaceVariantEntry,
    build_col_index,
    check_case_variant,
)
from .parsing import (
    append_note,
    chunked,
    normalize_key,
    parse_blr_index,
    parse_uncertainty,
)
from .report import ImportReport

handler = RichHandler(
    rich_tracebacks=True,
    tracebacks_show_locals=True,
    markup=True,
    show_time=True,
    show_path=True,
)

logging.basicConfig(level=logging.INFO, format="%(message)s", handlers=[handler])
logger = logging.getLogger(__name__)


def main() -> None:
    start_time = time.monotonic()
    logger.info(f"Using database at {DB_STRING}")

    create_all(engine)

    if not EXCEL.exists():
        logger.warning(f"Excel file not found at {EXCEL}")
        return

    logger.info(f"Reading: {EXCEL}")
    wb = load_workbook(EXCEL, read_only=True, data_only=True)
    ws = wb.active
    logger.info(f"Worksheet: '{ws.title}' | rows (approx): {ws.max_row}")

    rows = ws.rows
    try:
        header_row = next(rows)
    except StopIteration:
        logger.error("Empty Excel file.")
        return

    headers = [
        str(c.value).strip() if c.value is not None else None for c in header_row
    ]

    logger.info(f"Columns: {[h for h in headers if h]}")

    import_report = ImportReport()

    # Reconstructions first, so each reflex is created with its blr_entry_id.
    import_blr(import_report)
    blr_ids = reconstruction_ids()

    file_name = EXCEL.name
    sheet_name = ws.title
    col_index = build_col_index(headers)

    # Check for missing expected columns. Collapse internal whitespace so that
    # e.g. the "Consultant  2" (double space) header still matches "Consultant 2".
    def _collapse_ws(s: str) -> str:
        return " ".join(s.split())

    header_set = {_collapse_ws(h) for h in headers if h}
    for col in EXPECTED_COLUMNS:
        if _collapse_ws(col) not in header_set:
            import_report.add(
                MissingColumnEntry(
                    loc=SourceLocation(file_name, sheet_name, None, col, None),
                    note="Verwachte kolom niet aanwezig in het Excel-bestand.",
                )
            )
            logger.warning(f"Missing expected column: '{col}'")

    def row_to_dict(rc) -> dict:
        return {h: c.value for h, c in zip(headers, rc, strict=False) if h is not None}

    # -----------------------------------------------------------------------
    # In-memory caches (key -> ORM object)
    # -----------------------------------------------------------------------
    varieties_cache: dict[tuple, Variety] = {}
    countries_cache: dict[tuple, Country] = {}
    regions_cache: dict[tuple, Region] = {}
    places_cache: dict[tuple, Place] = {}
    locations_cache: dict[tuple, Location] = {}
    sources_cache: dict[tuple, Source] = {}
    projects_cache: dict[tuple, Project] = {}
    years_cache: dict[tuple, CollectionYear] = {}
    data_collectors_cache: dict[tuple, DataCollector] = {}
    consultants_cache: dict[tuple, Consultant] = {}
    fieldworks_cache: dict[tuple, Fieldwork] = {}

    # Normalised caches for case-variant detection (normalised_key -> original_key)
    varieties_norm: dict[tuple, tuple] = {}
    countries_norm: dict[tuple, tuple] = {}
    regions_norm: dict[tuple, tuple] = {}
    data_collectors_norm: dict[tuple, tuple] = {}
    consultants_norm: dict[tuple, tuple] = {}
    sources_norm: dict[tuple, tuple] = {}

    # (location_id, description) -> (lat, lon, excel_row)
    coord_seen: dict[tuple[int, str], tuple[float, float, int]] = {}
    # beginning_id -> first excel_row
    seen_beginning_ids: dict[int, int] = {}
    # (term, variety_id) -> first excel_row
    seen_term_variety: dict[tuple[str, Any], int] = {}

    total_processed = 0
    # excel_row -> reflex id, so report entries can link to the admin
    row_reflex: dict[int, int] = {}

    with Session(engine, expire_on_commit=False) as session:
        # -------------------------------------------------------------------
        # Pre-load caches from existing DB rows
        # -------------------------------------------------------------------
        logger.info("Pre-loading caches from database...")

        def _preload(model, key_fn, cache, norm=None, guard=None):
            for x in session.exec(select(model)).all():
                if guard and not guard(x):
                    continue
                key = key_fn(x)
                cache[key] = x
                if norm is not None:
                    norm.setdefault(normalize_key(key), key)

        _preload(
            Variety,
            lambda v: (
                v.name_english,
                v.name_french,
                v.name_bantu,
                v.name_bastin,
                v.branch,
                v.bastin_index,
                v.guthrie_index,
                v.guthrie_plus,
            ),
            varieties_cache,
            varieties_norm,
        )
        _preload(Country, lambda c: (c.name,), countries_cache, countries_norm,
                 guard=lambda c: c.name)
        _preload(Region, lambda r: (r.name, r.country_id), regions_cache, regions_norm,
                 guard=lambda r: r.name)
        _preload(Place, lambda p: (p.name, p.region_id), places_cache,
                 guard=lambda p: p.name)
        _preload(Location, lambda loc: (loc.place_id, loc.region_id, loc.country_id),
                 locations_cache)
        _preload(Source, lambda s: (s.author_year, s.title), sources_cache, sources_norm)
        _preload(Project, lambda p: (p.name,), projects_cache, guard=lambda p: p.name)
        _preload(CollectionYear, lambda cy: (cy.value, cy.status), years_cache)
        _preload(DataCollector, lambda dc: (dc.name,), data_collectors_cache,
                 data_collectors_norm, guard=lambda dc: dc.name)
        _preload(Consultant, lambda c: (c.name,), consultants_cache, consultants_norm,
                 guard=lambda c: c.name)
        # Bulk-load link tables once instead of N+1 queries per fieldwork.
        fw_dc_map: dict[int, list[int]] = {}
        for lnk in session.exec(select(FieldworkDataCollector)).all():
            fw_dc_map.setdefault(lnk.fieldwork_id, []).append(lnk.data_collector_id)
        fw_con_map: dict[int, list[int]] = {}
        for lnk in session.exec(select(FieldworkConsultant)).all():
            fw_con_map.setdefault(lnk.fieldwork_id, []).append(lnk.consultant_id)
        for fw in session.exec(select(Fieldwork)).all():
            dc_ids = tuple(sorted(fw_dc_map.get(fw.id, [])))
            con_ids = tuple(sorted(fw_con_map.get(fw.id, [])))
            k = (
                fw.project_id,
                fw.collection_year_id,
                fw.archive,
                fw.data_entered_by,
                dc_ids,
                con_ids,
            )
            fieldworks_cache[k] = fw

        logger.info("Caches pre-loaded.")

        # -------------------------------------------------------------------
        # Row processing
        # -------------------------------------------------------------------
        for batch in chunked(rows, 5000):
            # Deferred per-batch buffers: avoid one flush-with-RETURNING per row.
            batch_reflexes: list[Reflex] = []
            batch_source_links: list[tuple[Reflex, list[dict]]] = []
            batch_rows: list[int] = []
            # (reflex, BLR number, location) of numbers without a reconstruction; reported after the flush.
            batch_unresolved: list[tuple[Reflex, int, SourceLocation]] = []
            try:
                for i, row_cells in enumerate(batch):
                    excel_row = total_processed + i + 2
                    row = row_to_dict(row_cells)
                    # Fully empty row (trailing Excel padding): skip silently,
                    # this is not a data problem worth reporting.
                    if not any(
                        v is not None and str(v).strip() for v in row.values()
                    ):
                        continue
                    ctx = RowContext(
                        file_name, sheet_name, excel_row, col_index, import_report, row
                    )
                    try:
                        # -- Whitespace variant detection --
                        for _col in WHITESPACE_CHECK_FIELDS:
                            _raw = row.get(_col)
                            if (
                                _raw is not None
                                and isinstance(_raw, str)
                                and _raw != _raw.strip()
                            ):
                                import_report.add(
                                    WhitespaceVariantEntry(
                                        loc=ctx.loc(_col),
                                        note="Waarde bevat overbodige spaties — automatisch gestript.",
                                        raw_value=repr(_raw),
                                        stripped_value=_raw.strip(),
                                    )
                                )

                        # 1. Variety
                        v_key = (
                            ctx.val("Variety (English)"),
                            ctx.val("variety (French)"),
                            ctx.val("Variety (Bantu)"),
                            ctx.val("Variety as named in Bastin"),
                            ctx.val("Branch"),
                            ctx.val("Bastin Index"),
                            ctx.val("Guthrie Index"),
                            ctx.val("Guthrie PLUS"),
                            ctx.val("Branch (detailed)"),
                        )
                        check_case_variant(
                            varieties_norm,
                            v_key,
                            "Variety",
                            ctx,
                            column="Variety (English)",
                            term=ctx.val("Term"),
                        )
                        variety = get_or_create(
                            varieties_cache,
                            v_key,
                            lambda k=v_key: Variety(
                                name_english=k[0],
                                name_french=k[1],
                                name_bantu=k[2],
                                name_bastin=k[3],
                                branch=k[4],
                                bastin_index=k[5],
                                guthrie_index=k[6],
                                guthrie_plus=k[7],
                                branch_detailed=k[8],
                            ),
                            session,
                        )

                        # 2. Location hierarchy (clan excluded per project decision)
                        country_val = ctx.val("Country")
                        if country_val:
                            check_case_variant(
                                countries_norm,
                                (country_val,),
                                "Country",
                                ctx,
                                column="Country",
                            )
                        country = (
                            get_or_create(
                                countries_cache,
                                (country_val,),
                                lambda n=country_val: Country(name=n),
                                session,
                            )
                            if country_val
                            else None
                        )

                        region_val = ctx.val("Region (fieldwork)")
                        r_key = (region_val, country.id if country else None)
                        if region_val:
                            check_case_variant(
                                regions_norm,
                                r_key,
                                "Region",
                                ctx,
                                column="Region (fieldwork)",
                            )
                        region = (
                            get_or_create(
                                regions_cache,
                                r_key,
                                lambda k=r_key: Region(name=k[0], country_id=k[1]),
                                session,
                            )
                            if region_val
                            else None
                        )

                        place_val = ctx.val("Place")
                        p_key = (place_val, region.id if region else None)
                        place = (
                            get_or_create(
                                places_cache,
                                p_key,
                                lambda k=p_key: Place(name=k[0], region_id=k[1]),
                                session,
                            )
                            if place_val
                            else None
                        )

                        loc_key = (
                            place.id if place else None,
                            region.id if region else None,
                            country.id if country else None,
                        )
                        location = get_or_create(
                            locations_cache,
                            loc_key,
                            lambda k=loc_key: Location(
                                place_id=k[0], region_id=k[1], country_id=k[2]
                            ),
                            session,
                        )

                        # 3. Coordinates (BK/BL primary; BM/BN as fallback)
                        lat_def = ctx.float("Latitude")
                        lon_def = ctx.float("Longitude")
                        lat_src = ctx.float("Latitude OR source")
                        lon_src = ctx.float("Longitude OR source")
                        coord_source = ctx.val("Source of coordinates")

                        if (
                            (lat_def is None or lon_def is None)
                            and lat_src is not None
                            and lon_src is not None
                        ):
                            lat_def, lon_def = lat_src, lon_src
                            import_report.add(
                                FallbackUsedEntry(
                                    loc=ctx.loc("Latitude"),
                                    note="BK/BL leeg; gekopieerd van BM/BN (OR source)",
                                    lat_used=lat_def,
                                    lon_used=lon_def,
                                )
                            )

                        def _add_coord(
                            lat: float,
                            lon: float,
                            desc: str,
                            *,
                            _loc=location,
                            _ctx=ctx,
                            _src=coord_source,
                        ) -> None:
                            if _loc is None:
                                return
                            tk = (_loc.id, desc)
                            if tk not in coord_seen:
                                session.add(
                                    Coordinate(
                                        location_id=_loc.id,
                                        latitude=lat,
                                        longitude=lon,
                                        description=desc,
                                        coordinate_source=_src,
                                    )
                                )
                                coord_seen[tk] = (lat, lon, _ctx.row)
                            else:
                                ex_lat, ex_lon, ex_row = coord_seen[tk]
                                if lat != ex_lat or lon != ex_lon:
                                    import_report.add(
                                        CoordinateConflictEntry(
                                            loc=_ctx.loc("Latitude"),
                                            note="Zelfde locatie met afwijkende coördinaten.",
                                            coord_type=desc,
                                            existing_lat=ex_lat,
                                            existing_lon=ex_lon,
                                            new_lat=lat,
                                            new_lon=lon,
                                            term=_ctx.val("Term"),
                                            variety_english=_ctx.val(
                                                "Variety (English)"
                                            ),
                                            country=_ctx.val("Country"),
                                            place=_ctx.val("Place"),
                                            conflicting_row=ex_row,
                                        )
                                    )

                        if lat_def is not None and lon_def is not None:
                            _add_coord(lat_def, lon_def, "default")
                        if lat_src is not None and lon_src is not None:
                            _add_coord(lat_src, lon_src, "original")

                        # 4. Fieldwork

                        # 4a. Project
                        project_val = ctx.val("Project")
                        project = (
                            get_or_create(
                                projects_cache,
                                (project_val,),
                                lambda n=project_val: Project(name=n),
                                session,
                            )
                            if project_val
                            else None
                        )

                        # 4b. Collection year
                        f_year_raw = ctx.val("Year of collection")
                        collection_year = None
                        if f_year_raw is not None:
                            f_year_val, f_unc = parse_uncertainty(f_year_raw)
                            if f_year_val is None:
                                year_status = "unknown"
                            elif f_unc:
                                year_status = "uncertain"
                            else:
                                year_status = "known"

                            if f_unc:
                                import_report.add(
                                    UncertaintyEntry(
                                        loc=ctx.loc("Year of collection"),
                                        note="Jaar onzeker; controleer op alternatief jaar indien aanwezig.",
                                        original_value=f_year_raw,
                                        cleaned_value=f_year_val,
                                        status=year_status,
                                    )
                                )

                            year_key = (f_year_val, year_status)
                            collection_year = get_or_create(
                                years_cache,
                                year_key,
                                lambda k=year_key: CollectionYear(
                                    value=k[0], status=k[1]
                                ),
                                session,
                            )

                        # 4c. Data collectors (dynamic detection)
                        dc_names = sorted(
                            # set: the same person can be listed in two columns
                            {
                                ctx.val(h)
                                for h in headers
                                if h and COLLECTOR_PATTERN.match(h)
                            }
                            - {None}
                        )
                        dc_objects = []
                        for name in dc_names:
                            check_case_variant(
                                data_collectors_norm,
                                (name,),
                                "DataCollector",
                                ctx,
                            )
                            dc_objects.append(
                                get_or_create(
                                    data_collectors_cache,
                                    (name,),
                                    lambda n=name: DataCollector(name=n),
                                    session,
                                )
                            )

                        # 4d. Consultants (dynamic detection)
                        con_names = sorted(
                            # set: the same person can be listed in two columns
                            {
                                ctx.val(h)
                                for h in headers
                                if h and CONSULTANT_PATTERN.match(h)
                            }
                            - {None}
                        )
                        con_objects = []
                        for name in con_names:
                            check_case_variant(
                                consultants_norm,
                                (name,),
                                "Consultant",
                                ctx,
                            )
                            con_objects.append(
                                get_or_create(
                                    consultants_cache,
                                    (name,),
                                    lambda n=name: Consultant(name=n),
                                    session,
                                )
                            )

                        # 4e. Fieldwork
                        archive_val = ctx.val("Archive")
                        entered_by = ctx.val("Data entered by")
                        dc_id_tuple = tuple(sorted(dc.id for dc in dc_objects))
                        con_id_tuple = tuple(sorted(c.id for c in con_objects))

                        f_key = (
                            project.id if project else None,
                            collection_year.id if collection_year else None,
                            archive_val,
                            entered_by,
                            dc_id_tuple,
                            con_id_tuple,
                        )
                        fieldwork = None
                        if any(k is not None for k in f_key):
                            if f_key in fieldworks_cache:
                                fieldwork = fieldworks_cache[f_key]
                            else:
                                fieldwork = Fieldwork(
                                    project_id=project.id if project else None,
                                    collection_year_id=collection_year.id
                                    if collection_year
                                    else None,
                                    archive=archive_val,
                                    data_entered_by=entered_by,
                                )
                                session.add(fieldwork)
                                session.flush()
                                for dc in dc_objects:
                                    session.add(
                                        FieldworkDataCollector(
                                            fieldwork_id=fieldwork.id,
                                            data_collector_id=dc.id,
                                        )
                                    )
                                for con in con_objects:
                                    session.add(
                                        FieldworkConsultant(
                                            fieldwork_id=fieldwork.id,
                                            consultant_id=con.id,
                                        )
                                    )
                                fieldworks_cache[f_key] = fieldwork

                        # 5. Sources
                        source_objects = []
                        for auth_col, title_col, page_col, role in SOURCE_GROUPS:
                            auth_raw = ctx.val(auth_col)
                            if not auth_raw:
                                continue

                            title_val = ctx.val(title_col) if title_col else None
                            page_val = ctx.val(page_col) if page_col else None

                            # Basic validation: if title column exists, both should be present
                            if title_col and bool(auth_raw) != bool(title_val):
                                import_report.add(
                                    SourceIncompleteEntry(
                                        loc=ctx.loc(auth_col),
                                        note="Bron heeft wel auteur/jaar maar geen titel, of omgekeerd.",
                                        author_year_column=auth_col,
                                        author_year_value=auth_raw or "(leeg)",
                                        title_column=title_col,
                                        title_value=title_val or "(leeg)",
                                    )
                                )

                            auth_val, unc = parse_uncertainty(auth_raw)
                            if unc:
                                import_report.add(
                                    UncertaintyEntry(
                                        loc=ctx.loc(auth_col),
                                        note="Onzeker jaar; controleer op alternatief indien aanwezig.",
                                        original_value=auth_raw,
                                        cleaned_value=auth_val,
                                    )
                                )

                            s_key = (auth_val, title_val)
                            check_case_variant(
                                sources_norm,
                                s_key,
                                "Source",
                                ctx,
                                column=auth_col,
                            )
                            source = get_or_create(
                                sources_cache,
                                s_key,
                                lambda k=s_key, u=unc: Source(
                                    author_year=k[0], year_uncertain=u, title=k[1]
                                ),
                                session,
                            )
                            # Only "????" without a title: nothing to link
                            # (already reported as uncertainty above).
                            if source is None:
                                continue

                            # Determine order: increment for the same role within this record
                            order = 1 + sum(
                                1 for s in source_objects if s["role"] == role
                            )
                            source_objects.append(
                                {
                                    "source_id": source.id,
                                    "role": role,
                                    "order": order,
                                    "pages": page_val,
                                }
                            )

                        # 6. Reflex
                        original_id = ctx.int("ID 20250825")
                        beginning_id = ctx.int("ID beginning")

                        if beginning_id is not None:
                            if beginning_id in seen_beginning_ids:
                                import_report.add(
                                    DuplicateBeginningIdEntry(
                                        loc=ctx.loc("ID beginning"),
                                        note="Dubbel ID beginning — ID beginning leeggemaakt; nieuw ID toekennen.",
                                        beginning_id=beginning_id,
                                        first_seen_row=seen_beginning_ids[
                                            beginning_id
                                        ],
                                        term=ctx.val("Term"),
                                        variety_english=ctx.val("Variety (English)"),
                                    )
                                )
                                # Split rows (e.g. "kíla ou bóondo") share the ID:
                                # import them, but without the unique beginning_id.
                                beginning_id = None
                            else:
                                seen_beginning_ids[beginning_id] = excel_row

                        term = ctx.val("Term")
                        if not term:
                            import_report.add(
                                SkippedNoTermEntry(
                                    loc=ctx.loc("Term"),
                                    note="Rij geskipt: geen waarde in kolom 'Term'.",
                                    variety_english=v_key[0],
                                    guthrie_index=v_key[6],
                                    country=ctx.val("Country"),
                                )
                            )
                            continue

                        # Suspicious duplicate: same term + variety combination
                        tv_key = (term, variety.id if variety else None)
                        if tv_key in seen_term_variety:
                            import_report.add(
                                SuspiciousDuplicateEntry(
                                    loc=ctx.loc("Term"),
                                    note="Zelfde term + variëteit al eerder gezien — mogelijke duplicaat.",
                                    first_seen_row=seen_term_variety[tv_key],
                                    term=term,
                                    variety_english=v_key[0],
                                    guthrie_index=v_key[6],
                                )
                            )
                        else:
                            seen_term_variety[tv_key] = excel_row

                        blr_raw = ctx.val("BLR Index")
                        blr_index, blr_text, used_ipv = parse_blr_index(blr_raw)
                        blr_entry_id = blr_ids.get(int(blr_index)) if blr_index else None
                        # Cell text that is not a plain linked number stays in the note.
                        reconstruction_note = append_note(
                            None, blr_text or (None if blr_entry_id else blr_index)
                        )
                        if used_ipv:
                            import_report.add(
                                BlrIpvResolvedEntry(
                                    loc=ctx.loc("BLR Index"),
                                    note="'ipv' in BLR index: eerste nummer gebruikt en gelinkt.",
                                    blr_index_raw=blr_raw,
                                    blr_index_used=blr_index,
                                    term=term,
                                    variety_english=v_key[0],
                                )
                            )
                        elif blr_text is not None:
                            import_report.add(
                                BlrTextToNoteEntry(
                                    loc=ctx.loc("BLR Index"),
                                    note="BLR index is geen getal: niet gelinkt, bewaard in reconstruction_note.",
                                    blr_index_raw=blr_raw,
                                    term=term,
                                    variety_english=v_key[0],
                                )
                            )

                        reflex = Reflex(
                            original_id=original_id,
                            beginning_id=beginning_id,
                            term=term,
                            term_morph_analysis=ctx.val(
                                "Term with Morphological Analysis"
                            ),
                            translation=ctx.val("Translation"),
                            meaning=ctx.val("Meaning"),
                            scientific_name=ctx.val("Scientific name"),
                            semantic_field=ctx.val("Semantic field"),
                            part_of_speech=ctx.val("Part of Speech"),
                            reconstruction_note=reconstruction_note,
                            blr_entry_id=blr_entry_id,
                            ipa=ctx.val("IPA"),
                            notes_on_ipa=ctx.val("Notes on IPA"),
                            notes=ctx.val("Opmerking"),
                            variety_id=variety.id if variety else None,
                            location_id=location.id if location else None,
                            fieldwork_id=fieldwork.id if fieldwork else None,
                        )
                        # Defer add/flush — done once per batch below.
                        batch_reflexes.append(reflex)
                        batch_rows.append(excel_row)
                        batch_source_links.append((reflex, source_objects))
                        if blr_index and blr_entry_id is None:
                            batch_unresolved.append((reflex, int(blr_index), ctx.loc("BLR Index")))

                    except Exception as e:
                        logger.error(f"Error on row {excel_row}: {e}")
                        import_report.add(
                            CriticalErrorEntry(
                                loc=ctx.loc(),
                                note="Onverwachte fout tijdens verwerking van rij.",
                                error_message=str(e),
                            )
                        )
                        continue

                # Bulk-insert all reflexes for this batch → one round trip for IDs,
                # then bulk-insert their SourceEntries with resolved entry_ids.
                if batch_reflexes:
                    session.add_all(batch_reflexes)
                    session.flush()
                    for reflex, idx, loc in batch_unresolved:
                        import_report.add(
                            BlrUnresolvedLinkEntry(
                                loc=loc,
                                note="Geen reconstructie met deze BLR index: niet gelinkt, bewaard in reconstruction_note.",
                                reflex_id=reflex.id,
                                blr_index_raw=idx,
                            )
                        )
                    for reflex, source_objects in batch_source_links:
                        for s_data in source_objects:
                            session.add(
                                SourceEntries(
                                    entry_id=reflex.id,
                                    source_id=s_data["source_id"],
                                    role=s_data["role"],
                                    order=s_data["order"],
                                    pages=s_data["pages"],
                                )
                            )

                session.commit()
                row_reflex.update(
                    zip(batch_rows, (r.id for r in batch_reflexes), strict=True)
                )
                total_processed += len(batch)
                logger.info(
                    f"Committed batch — total rows processed: {total_processed}"
                )

            except Exception as e:
                logger.error(f"Batch error: {e}")
                session.rollback()
                continue

    rows_path = DATA_ROOT / "import_rows.json"
    rows_path.write_text(json.dumps(row_reflex))
    logger.info(f"Row map ({len(row_reflex)} reflexes) saved to {rows_path}")

    elapsed = time.monotonic() - start_time
    logger.info(f"Import complete in {elapsed:.1f}s.")

    # Validation summary in terminal, coloured per severity level
    _level_style = {"error": "bold red", "warning": "yellow", "note": "dim"}
    summary = import_report.summary()
    if summary:
        logger.warning("Validation summary:")
        for category, count in sorted(
            summary.items(),
            key=lambda x: (LEVEL_ORDER[CATEGORY_CLASSES[x[0]].level], -x[1]),
        ):
            level = CATEGORY_CLASSES[category].level
            style = _level_style[level]
            logger.warning(
                f"  [{style}]{level:<7}[/] {category}: {count} issue(s)"
            )
    else:
        logger.info("Validation: geen problemen gevonden.")

    import_report.save(
        str(DATA_ROOT / "import_validation.xlsx"),
        str(DATA_ROOT / "import_validation.json"),
    )
    logger.info(
        f"Validation report saved to {DATA_ROOT / 'import_validation.xlsx'} "
        f"and {DATA_ROOT / 'import_validation.json'}"
    )


if __name__ == "__main__":
    main()
