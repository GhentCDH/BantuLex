import json
from dataclasses import asdict
from typing import ClassVar

from .entries import (
    CATEGORY_CLASSES,
    FIELD_HEADERS,
    LEVEL_ERROR,
    LEVEL_NOTE,
    LEVEL_WARNING,
    LOCATION_HEADERS,
    ReportEntry,
    entry_fields,
)

# Styling per severity level (Excel fills and sheet tab colours).
_LEVEL_FILL = {
    LEVEL_ERROR: "FFC7CE",  # red
    LEVEL_WARNING: "FFEB9C",  # orange/yellow
    LEVEL_NOTE: "D9D9D9",  # grey
}
_LEVEL_FONT = {
    LEVEL_ERROR: "9C0006",
    LEVEL_WARNING: "9C6500",
    LEVEL_NOTE: "595959",
}

# Overview columns that are filled from matching entry fields when present.
_OVERVIEW_ORIGINAL = ("original_value", "raw_value")
_OVERVIEW_CLEANED = ("cleaned_value", "stripped_value")


class ImportReport:
    CATEGORIES: ClassVar[list[str]] = list(CATEGORY_CLASSES)

    def __init__(self) -> None:
        self.entries: dict[str, list[ReportEntry]] = {c: [] for c in self.CATEGORIES}
        self._seen: set[ReportEntry] = set()

    def add(self, entry: ReportEntry) -> None:
        # Frozen dataclasses are hashable: identical entries (same location,
        # values and note) are the same problem reported twice — keep one.
        try:
            if entry in self._seen:
                return
            self._seen.add(entry)
        except TypeError:
            pass  # unhashable field value — keep without dedup
        self.entries[entry.category].append(entry)

    def summary(self) -> dict[str, int]:
        """Return a dict of category -> count for non-empty categories."""
        return {c: len(v) for c, v in self.entries.items() if v}

    def save(self, xlsx_path: str, json_path: str) -> None:
        self._save_xlsx(xlsx_path)
        self._save_json(json_path)

    # -- Excel ------------------------------------------------------------

    def _save_xlsx(self, path: str) -> None:
        from openpyxl import Workbook

        # write_only streams rows to disk instead of building the full
        # workbook in memory — much faster for large reports.
        wb = Workbook(write_only=True)

        self._write_overview(wb)

        for category in self.CATEGORIES:
            entries = self.entries.get(category, [])
            cls = CATEGORY_CLASSES[category]
            ws = wb.create_sheet(title=category.upper()[:31])
            ws.sheet_properties.tabColor = _LEVEL_FILL[cls.level]
            if not entries:
                ws.append(["(no entries)"])
                continue
            cat_fields = entry_fields(cls)
            headers = (
                [h for _, h in LOCATION_HEADERS]
                + [FIELD_HEADERS.get(f, f) for f in cat_fields]
                + ["Note"]
            )
            ws.append(headers)
            for entry in sorted(entries, key=lambda e: e.loc.row or 0):
                ws.append(
                    [_cell(getattr(entry.loc, f)) for f, _ in LOCATION_HEADERS]
                    + [_cell(getattr(entry, f)) for f in cat_fields]
                    + [entry.note]
                )

        wb.save(path)

    def _write_overview(self, wb) -> None:
        """First sheet: warnings and errors merged, grouped per source row.

        Notes (harmless, automatically handled issues) are excluded here;
        they remain available in their own category sheets and in the JSON.
        The row number is printed only on the first entry of each row group;
        blank means "applies to the last-mentioned row above". Row-less
        entries (missing columns, BLR links) are grouped at the top.
        """
        from openpyxl.cell import WriteOnlyCell
        from openpyxl.styles import Font, PatternFill

        ws = wb.create_sheet(title="OVERVIEW")
        ws.append(
            [
                "Row",
                "Level",
                "Category",
                "Sheet",
                "Cell",
                "Column",
                "Original Value",
                "Cleaned Value",
                "Details",
                "Note",
            ]
        )

        merged = [
            (category, entry)
            for category in self.CATEGORIES
            if CATEGORY_CLASSES[category].level != LEVEL_NOTE
            for entry in self.entries.get(category, [])
        ]
        merged.sort(key=lambda ce: (ce[1].loc.row is not None, ce[1].loc.row or 0))

        fills = {lv: PatternFill("solid", fgColor=c) for lv, c in _LEVEL_FILL.items()}
        fonts = {lv: Font(color=c) for lv, c in _LEVEL_FONT.items()}

        last_row: int | None | str = "sentinel"
        for category, entry in merged:
            row = entry.loc.row
            if row != last_row:
                row_label = row if row is not None else "(geen rij)"
                last_row = row
            else:
                row_label = ""
            original = _first_field(entry, _OVERVIEW_ORIGINAL)
            cleaned = _first_field(entry, _OVERVIEW_CLEANED)
            details = "; ".join(
                f"{FIELD_HEADERS.get(f, f)}={getattr(entry, f)}"
                for f in entry_fields(type(entry))
                if f not in _OVERVIEW_ORIGINAL + _OVERVIEW_CLEANED
                and getattr(entry, f) is not None
            )
            level = type(entry).level
            level_cell = WriteOnlyCell(ws, value=level)
            level_cell.fill = fills[level]
            level_cell.font = fonts[level]
            ws.append(
                [
                    row_label,
                    level_cell,
                    category,
                    entry.loc.sheet,
                    _cell(entry.loc.cell),
                    _cell(entry.loc.column),
                    _cell(original),
                    _cell(cleaned),
                    details,
                    entry.note,
                ]
            )

    # -- JSON -------------------------------------------------------------

    def _save_json(self, path: str) -> None:
        data = {}
        for category in self.CATEGORIES:
            out = []
            for entry in sorted(
                self.entries.get(category, []), key=lambda e: e.loc.row or 0
            ):
                d = asdict(entry)
                d["location"] = d.pop("loc")
                d["level"] = type(entry).level
                out.append(d)
            data[category] = out
        with open(path, "w", encoding="utf-8") as f:
            json.dump(
                data, f, ensure_ascii=False, separators=(",", ":"), default=str
            )


def _cell(value):
    """Excel-safe cell value ('' for None, str() for exotic types)."""
    if value is None:
        return ""
    if isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def _first_field(entry: ReportEntry, names: tuple[str, ...]):
    for name in names:
        if hasattr(entry, name):
            val = getattr(entry, name)
            if val is not None:
                return val
    return None
